#!/usr/bin/env python3
import json
import re
import os
import argparse
import urllib.request
import urllib.error
import urllib.parse
import ssl
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime

# Setup paths
base_dir = os.path.dirname(os.path.abspath(__file__))
companies_path = os.path.join(base_dir, "companies.json")
manual_jobs_path = os.path.join(base_dir, "manual_jobs.json")
output_dir = os.path.join(base_dir, "Web3_Jobs_Research")
dashboard_path = os.path.join(base_dir, "jobs_dashboard.html")

# 35 Categories mapping
categories = {
    1: ("BD / Sales / Partnerships Jobs", r"\b(business development|bd|sales|partnership|partner|alliance|ecosystem|commercial|growth|account executive|client relationship|relationship manager|key account)\b"),
    2: ("Developers Jobs", r"\b(developer|engineer|programmer|architect|coder|solidity|rust|golang|backend|frontend|fullstack|full-stack|smart contract|protocol|devops|infrastructure|evm|android|ios|kotlin|swift|compiler|embedded|systems engineer)\b"),
    3: ("Marketing Jobs", r"\b(marketing|social media|content|brand|pr |public relations|communications|seo|copywriter|growth marketer|digital marketing|marketing manager|creative director|marketing specialist|copy writer|newsletter)\b"),
    4: ("Crypto Jobs in Indonesia - Bali", r"\b(indonesia|bali|jakarta)\b"),
    5: ("C-Level / Executive Jobs", r"\b(ceo|cto|coo|cfo|cmo|cpo|vice president|vp|president|executive director|managing director|chief of staff|general manager|country manager|managing partner|founder|co-founder)\b"),
    6: ("Trading & Hedge Funds Jobs", r"\b(trader|trading|hedge fund|portfolio manager|market maker|quant|quantitative|hft|arbitrage|execution trader|algorithmic trader|treasury trader|liquidity provider)\b"),
    7: ("Singapore / APAC Crypto Jobs", r"\b(singapore|sg|apac|asia|sydney|australia|vietnam|thailand|korea|seoul)\b"),
    8: ("Hong Kong Crypto Jobs", r"\b(hong kong|hk)\b"),
    9: ("Compliance / Legal Jobs", r"\b(compliance|legal|counsel|aml|kyc|regulatory|mlro|risk manager|sanctions|lawyer|audit|attorney|policy|regulatory affairs|general counsel)\b"),
    10: ("Stablecoin Jobs", r"\b(stablecoin|payments|reserve|liquidity manager|payment operations|tether|circle|ripple|paxos|ethena|cbdc|fiat)\b"),
    11: ("Dubai Jobs", r"\b(dubai|uae|united arab emirates|abu dhabi|mena|middle east|riyadh|saudi)\b"),
    12: ("Signal Board | all crypto roles", r"."),
    13: ("Tech Job board - general jobs in tech", r"\b(software engineer|product manager|data scientist|devops|engineering manager|full stack|sre|mobile engineer|qa engineer|program manager|cloud architect|machine learning|ux researcher|platform engineer|it support|helpdesk)\b"),
    14: ("Crypto Research & Analytics Jobs", r"\b(research|researcher|tokenomics|on-chain analyst|fundamental analyst|analytics|data analyst|data scientist|dune|bi developer|sql|defi analyst)\b"),
    15: ("For Founders: Hiring Tips by Jaya Talent", r"(?!)"),
    16: ("KOL Jobs", r"\b(kol|influencer|creator|ambassador|affiliate|tiktok|youtube|creator partnerships|influencer relations|key opinion leader)\b"),
    17: ("Japan Digital Assets Jobs", r"\b(japan|tokyo)\b"),
    18: ("Private Equity Jobs", r"\b(private equity|pe associate|pe analyst|lbo|m&a|mergers|buyout)\b"),
    19: ("HR & Talent Acquisition Jobs", r"\b(recruiter|talent|hr|human resources|people ops|people operations|sourcing specialist|headhunter|compensation|hiring manager|talent acquisition)\b"),
    20: ("Product & Project Jobs", r"\b(product manager|product owner|project manager|scrum master|program manager|product analyst|pm|head of product|vp of product|technical product manager)\b"),
    21: ("Designer Jobs", r"\b(design|ui/ux|ux/ui|graphic|visual|motion|creative|figma|brand designer|art director|illustrator|3d designer|web designer|product designer)\b"),
    22: ("VC / Investment Jobs", r"\b(investment|venture capital|vc|deal sourcing|portfolio analyst|venture analyst|general partner|principal|fund manager|investment associate|investment manager|venture partner)\b"),
    23: ("Finance · CFO · Treasury Jobs", r"\b(cfo|finance|financial|treasury|accountant|accounting|tax|controller|bookkeeper|fp&a|finance director|accounting manager|financial analyst)\b"),
    24: ("Customer Support · Trust & Safety", r"\b(support|customer success|helpdesk|trust & safety|trust and safety|fraud|moderator|dispute|customer support|support agent|customer care)\b"),
    25: ("Operations · COO · Chief of Staff", r"\b(operations|ops|coo|chief of staff|bizops|operational|operations manager|head of operations|business operations|ops manager|director of operations)\b"),
    26: ("SF Crypto Jobs", r"\b(san francisco|sf|bay area|oakland|san jose|california)\b"),
    27: ("NYC Crypto Jobs", r"\b(new york|nyc|brooklyn|manhattan|ny)\b"),
    28: ("London Crypto Jobs", r"\b(london|uk|united kingdom|great britain|england|britain)\b"),
    29: ("Swiss Crypto Jobs", r"\b(switzerland|zug|zurich|geneva|swiss|basel)\b"),
    30: ("India Crypto Jobs", r"\b(india|bangalore|bengaluru|mumbai|delhi|gurgaon|pune|hyderabad|chennai)\b"),
    31: ("Latam Crypto Jobs", r"\b(latam|latin america|brazil|mexico|argentina|colombia|sao paulo|mexico city|buenos aires|chile|peru|venezuela|ecuador)\b"),
    32: ("Community Manager Jobs", r"\b(community manager|discord|telegram|community lead|community moderator|community specialist|discord manager|telegram manager|community builder)\b"),
    33: ("DevRel Jobs", r"\b(devrel|developer relations|developer advocate|evangelist|technical writer|developer marketing|developer experience)\b"),
    34: ("AI x Web3 Jobs", r"\b(ai|artificial intelligence|machine learning|ml|llm|nlp|deep learning|agent|generative|openai|computer vision)\b"),
    35: ("Digital Assets Jobs", r"\b(digital asset|custody|asset management|tokenization|rwa|institutional digital assets|tokenised|tokenized)\b")
}

telegram_channels = {
    1: "t.me/web3bds", 2: "t.me/web3devs3", 3: "t.me/marketersjobs", 4: "t.me/jobsindonesiaa",
    5: "t.me/+FPxA5a5gdlQ3MTY6", 6: "t.me/AlphaHires", 7: "t.me/+n3TE8PhdosQ1YTcy", 8: "t.me/+cEVjTA5HfhhmZDNi",
    9: "t.me/compliancejobs", 10: "t.me/+gYai1mbLnIA5ODUy", 11: "t.me/dubaijobscrypto", 12: "t.me/+LuTusmxidno1MmQy",
    13: "t.me/TechJobsme", 14: "t.me/+FAi6UOcaLNUyMWMy", 15: "t.me/iamfuckingceo", 16: "t.me/+epk8zFJx3mM5ODE6",
    17: "t.me/japanjobss", 18: "t.me/privateequityjobs", 19: "t.me/web3headhunters", 20: "t.me/productjobscrypto",
    21: "t.me/web3designerjobs", 22: "t.me/web3vcjobs", 23: "t.me/+I28fCco1oV1jMjFi", 24: "t.me/+4vaSF4Pa9y1iMzJi",
    25: "t.me/+_SCikTdlBJU2MjAy", 26: "t.me/+KuA8mS1Rju1hMzcy", 27: "t.me/+ExouurtaF8JjNTli", 28: "t.me/+KPy0IcQLBnVlNGYy",
    29: "t.me/+TOZl7q7MmwhiNDc6", 30: "t.me/+6YoGHP7FLyRjZDgy", 31: "t.me/+EJRVUtvGKjAwODU6", 32: "t.me/+b3phRT_N51M5NWIy",
    33: "t.me/+EdL0eIFUvIw5Yjdi", 34: "t.me/+hrGsMucd6s5mN2Ey", 35: "t.me/+1v2qJM6gKLZiYWUy"
}

# HTTP fetch helper
def http_get(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'})
    context = ssl._create_unverified_context()
    with urllib.request.urlopen(req, timeout=8, context=context) as response:
        return response.read().decode('utf-8')

# Link verifier
def verify_url(url):
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/119.0.0.0 Safari/537.36'})
        context = ssl._create_unverified_context()
        with urllib.request.urlopen(req, timeout=10, context=context) as response:
            content = response.read().decode('utf-8', errors='ignore').lower()
            if "page not found" in content or "job is no longer available" in content or "404 not found" in content:
                return False
            return True
    except urllib.error.HTTPError as e:
        if e.code in [404, 410]:
            return False
        content = e.read().decode('utf-8', errors='ignore').lower()
        if "page not found" in content or "job is no longer available" in content or "404 not found" in content:
            return False
        return True
    except Exception:
        return False

# Token parser from URL
def parse_token(url, scraper_type):
    url = url.rstrip('/')
    if scraper_type == "GREENHOUSE":
        if url.endswith('/jobs'):
            url = url[:-5]
        return url.split('/')[-1]
    elif scraper_type == "LEVER":
        return url.split('/')[-1]
    elif scraper_type == "ASHBYHQ":
        return url.split('/')[-1]
    return None

def clean_job_data(job):
    title = job.get('title') or ''
    location = job.get('location') or ''
    # Clean HTML entities
    for ent, char in [('&amp;', '&'), ('&#x27;', "'"), ('&quot;', '"'), ('&lt;', '<'), ('&gt;', '>'), ('&#39;', "'")]:
        title = title.replace(ent, char)
        location = location.replace(ent, char)
    if title.endswith(' →'):
        title = title[:-2].strip()
    job['title'] = title
    job['location'] = location

    # Rewrite Worldcoin custom job links to direct employer links
    link = job.get('link') or ''
    if job.get('company', '').lower() == 'worldcoin' and 'ashbyhq.com' in link:
        uuid_part = link.rstrip('/').split('/')[-1]
        job['link'] = f"https://world.org/job/{uuid_part}"

    return job

def clean_greenhouse_link(link):
    if not link:
        return ''
    match = re.search(r'/([^/?]+)/(\d+)\?gh_jid=(\d+)', link)
    if match:
        path_name = match.group(1)
        jid = match.group(3)
        if "fireblocks.com" in link:
            return f"https://job-boards.greenhouse.io/embed/job_app?for=fireblocks&token={jid}"
        if "ripple.com" in link or "consensys.io" in link:
            parts = urllib.parse.urlparse(link)
            path = parts.path
            if not path.endswith('/'):
                path += '/'
            return urllib.parse.urlunparse((parts.scheme, parts.netloc, path, parts.params, f"gh_jid={jid}", parts.fragment))
        if "coinbase.com" in link:
            return link
        parts = urllib.parse.urlparse(link)
        new_path = parts.path.replace(f"/{path_name}/{jid}", f"/{path_name}/")
        new_path = re.sub(r'//+', '/', new_path)
        return urllib.parse.urlunparse((parts.scheme, parts.netloc, new_path, parts.params, f"gh_jid={jid}", parts.fragment))
    return link

# Fetchers for various boards
def fetch_greenhouse(company_name, token):
    url = f"https://boards-api.greenhouse.io/v1/boards/{token}/jobs"
    try:
        data = json.loads(http_get(url))
        jobs = data.get('jobs', [])
        return [clean_job_data({"company": company_name, "title": j.get('title'), "location": j.get('location', {}).get('name', 'Remote'), "link": clean_greenhouse_link(j.get('absolute_url'))}) for j in jobs]
    except Exception as e:
        return []

def fetch_lever(company_name, token):
    url = f"https://api.lever.co/v0/postings/{token}"
    try:
        data = json.loads(http_get(url))
        return [clean_job_data({"company": company_name, "title": j.get('text'), "location": j.get('categories', {}).get('location', 'Remote'), "link": j.get('hostedUrl')}) for j in data]
    except Exception as e:
        return []

def fetch_ashby(company_name, token):
    url = f"https://api.ashbyhq.com/posting-api/job-board/{token}"
    try:
        data = json.loads(http_get(url))
        jobs = data.get('jobs', [])
        return [clean_job_data({"company": company_name, "title": j.get('title'), "location": j.get('location', 'Remote'), "link": j.get('jobUrl')}) for j in jobs]
    except Exception as e:
        return []

def fetch_yzilabs(company_name, url):
    jobs = []
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            html = response.read().decode('utf-8')
            
        matches = re.findall(r'href="(/jobs/[a-f0-9\-]+)"[^>]*>(.*?)</a>', html, re.DOTALL)
        for path, inner in matches:
            clean = re.sub(r'<[^>]+>', ' ', inner)
            clean = re.sub(r'\s+', ' ', clean).strip()
            
            if ' · ' in clean:
                title, rest = clean.split(' · ', 1)
                title = title.strip()
                
                rest_parts = rest.strip().split(' ')
                company = rest_parts[0]
                
                location = "Remote"
                loc_match = re.search(r'Remote\s*(?:\([^)]+\))?', rest)
                if loc_match:
                    location = loc_match.group(0)
                elif len(rest_parts) > 1:
                    location = rest_parts[1]
                
                jobs.append(clean_job_data({
                    "company": company,
                    "title": title,
                    "location": location,
                    "link": f"https://talent.yzilabs.com{path}"
                }))
    except Exception:
        pass
    return jobs

def fetch_hashtagweb3(max_jobs=200):
    """Fetch from HashtagWeb3 public REST API - https://hashtagweb3.com/api/v1/jobs"""
    jobs = []
    cursor = None
    fetched = 0
    try:
        while fetched < max_jobs:
            url = 'https://hashtagweb3.com/api/v1/jobs?limit=50'
            if cursor:
                url += f'&cursor={urllib.parse.quote(cursor)}'
            req = urllib.request.Request(url, headers={
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
            })
            ctx = ssl._create_unverified_context()
            with urllib.request.urlopen(req, timeout=10, context=ctx) as res:
                data = json.loads(res.read().decode('utf-8'))
            page_jobs = data.get('data', [])
            if not page_jobs:
                break
            for j in page_jobs:
                jobs.append(clean_job_data({
                    'company': j.get('company', ''),
                    'title': j.get('title', ''),
                    'location': j.get('location', 'Remote'),
                    'link': j.get('link', ''),
                }))
            fetched += len(page_jobs)
            next_cursor = data.get('meta', {}).get('next_cursor')
            if not next_cursor:
                break
            cursor = next_cursor
    except Exception as e:
        print(f'[HashtagWeb3] fetch error: {e}')
    return jobs

def fetch_web3career(max_jobs=100):
    """Fetch from web3.career HTML listing - uses job page URL as the apply link."""
    jobs = []
    try:
        req = urllib.request.Request('https://web3.career/', headers={
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
            'Accept': 'text/html,application/xhtml+xml'
        })
        ctx = ssl._create_unverified_context()
        with urllib.request.urlopen(req, timeout=20, context=ctx) as res:
            html = res.read().decode('utf-8')
        # Each job row: <tr data-jobid=NNN onclick="tableTurboRowClick(event, '/slug/NNN')">...</tr>
        # Inside: alt="Company" or alt=Company (unquoted), <h2>Title</h2>, cell-location>Location<
        row_matches = re.findall(
            r'<tr\s+data-jobid=\d+\s+onclick="tableTurboRowClick\(event,\s*\'(/[^\']+)\'\)"[^>]*>(.*?)(?=<tr\s+data-jobid|\Z)',
            html, re.DOTALL
        )
        seen = set()
        for slug, body in row_matches:
            # Company: alt="Name" or alt=Name (unquoted)
            company_m = re.search(r'\balt=(?:"([^"]+)"|([A-Za-z][^\s>]*))', body)
            # Title: in <h2>
            title_m = re.search(r'<h2[^>]*>\s*([^<]+?)\s*</h2>', body)
            # Location: in cell-location
            loc_m = re.search(r'class=cell-location[^>]*>\s*([^<\s][^<]{1,40}?)\s*<', body)
            if not (company_m and title_m):
                continue
            company = (company_m.group(1) or company_m.group(2) or '').strip()
            title = title_m.group(1).strip()
            location = loc_m.group(1).strip() if loc_m else 'Remote'
            link = 'https://web3.career' + slug.strip()
            # Skip ads/metana sponsor links
            if 'metana' in slug or not company or len(title) < 4:
                continue
            if link not in seen:
                seen.add(link)
                jobs.append(clean_job_data({
                    'company': company,
                    'title': title,
                    'location': location or 'Remote',
                    'link': link,
                }))
            if len(jobs) >= max_jobs:
                break
    except Exception as e:
        print(f'[web3.career] fetch error: {e}')
    return jobs

# Interactive add CLI
def add_manual_job():
    print("--- Web3 Job Research: Add Custom Board Entry ---")
    company = input("Enter Company Name: ").strip()
    if not company:
        print("Company Name cannot be empty.")
        return
        
    title = input("Enter Role Title (e.g. Smart Contract Developer): ").strip()
    if not title:
        print("Title cannot be empty.")
        return
        
    location = input("Enter Location / Remote Status (e.g. London / Remote): ").strip()
    if not location:
        location = "Remote"
        
    link = input("Enter Direct Apply Link: ").strip()
    if not link:
        print("Apply Link cannot be empty.")
        return
        
    notes = input("Enter Notes (optional): ").strip()
    if not notes:
        notes = "Manually entered job posting from custom board."
        
    # Quick link test
    print("Verifying link connectivity...")
    if verify_url(link):
        print("✅ Link verified successfully.")
    else:
        print("⚠️ Warning: Link could not be verified (returned 404 or connection error).")
        proceed = input("Do you still want to add this link? (y/n): ").strip().lower()
        if proceed != 'y':
            print("Cancelled.")
            return

    new_job = {
        "company": company,
        "title": title,
        "location": location,
        "link": link,
        "date_found": datetime.now().strftime("%d %b %Y"),
        "notes": notes
    }
    
    # Load and save
    manual_jobs = []
    if os.path.exists(manual_jobs_path):
        try:
            with open(manual_jobs_path, 'r') as f:
                manual_jobs = json.load(f)
        except Exception:
            manual_jobs = []
            
    manual_jobs.append(new_job)
    
    with open(manual_jobs_path, 'w') as f:
        json.dump(manual_jobs, f, indent=4)
        
    print(f"🎉 Successfully added '{title}' at '{company}' to manual entries!")

# Load and fetch all jobs
def get_all_jobs():
    jobs_pool = []
    
    # 1. Load Manual Entries
    if os.path.exists(manual_jobs_path):
        try:
            with open(manual_jobs_path, 'r') as f:
                manual_jobs = json.load(f)
                print(f"Loaded {len(manual_jobs)} manual entries.")
                for mj in manual_jobs:
                    jobs_pool.append({
                        "company": mj.get('company'),
                        "title": mj.get('title'),
                        "location": mj.get('location'),
                        "link": mj.get('link'),
                        "notes": mj.get('notes'),
                        "is_manual": True
                    })
        except Exception as e:
            print(f"Error loading manual entries: {e}")

    # 2. Load and crawl companies
    if not os.path.exists(companies_path):
        print("Error: companies.json file not found in directory!")
        return jobs_pool
        
    with open(companies_path, 'r') as f:
        config = json.load(f)
    companies_list = config.get('companies', []) if isinstance(config, dict) else config
    active_companies = [c for c in companies_list if c.get('category') == 'crypto' and c.get('enabled', True)]
    
    # Greenhouse, Lever, Ashby, Yzilabs targets
    targets = []
    for c in active_companies:
        name = c.get('name')
        url = c.get('jobs_url')
        scraper = c.get('scraper')
        if scraper in ["GREENHOUSE", "LEVER", "ASHBYHQ"]:
            token = parse_token(url, scraper)
            if token:
                targets.append((name, scraper, token))
        elif scraper == "YZILABS":
            targets.append((name, scraper, url))
                
    print(f"Fetching active job listings for {len(targets)} Web3 companies via REST APIs...")
    
    # Thread pool for concurrency
    fetched_count = 0
    with ThreadPoolExecutor(max_workers=30) as executor:
        futures = {}
        for item in targets:
            if len(item) == 3 and item[1] in ["GREENHOUSE", "LEVER", "ASHBYHQ"]:
                name, scraper, token = item
                if scraper == "GREENHOUSE":
                    futures[executor.submit(fetch_greenhouse, name, token)] = (name, "Greenhouse")
                elif scraper == "LEVER":
                    futures[executor.submit(fetch_lever, name, token)] = (name, "Lever")
                elif scraper == "ASHBYHQ":
                    futures[executor.submit(fetch_ashby, name, token)] = (name, "Ashby")
            elif len(item) == 3 and item[1] == "YZILABS":
                name, scraper, url = item
                futures[executor.submit(fetch_yzilabs, name, url)] = (name, "YZi Labs")
                
        for future in as_completed(futures):
            name, board_type = futures[future]
            try:
                company_jobs = future.result()
                if company_jobs:
                    jobs_pool.extend(company_jobs)
                    fetched_count += len(company_jobs)
            except Exception:
                pass
                
    print(f"Completed fetching. Loaded {fetched_count} live listings from boards.")

    # Fetch HashtagWeb3 jobs
    print("Fetching jobs from HashtagWeb3...")
    hw3_jobs = fetch_hashtagweb3(max_jobs=200)
    jobs_pool.extend(hw3_jobs)
    print(f"HashtagWeb3: loaded {len(hw3_jobs)} jobs.")

    # Fetch web3.career jobs
    print("Fetching jobs from web3.career...")
    w3c_jobs = fetch_web3career(max_jobs=100)
    jobs_pool.extend(w3c_jobs)
    print(f"web3.career: loaded {len(w3c_jobs)} jobs.")
    
    # 3. Load pre-scraped jobs database for custom/other scrapers
    crypto_jobs_path = os.path.join(base_dir, "crypto_jobs.json")
    if os.path.exists(crypto_jobs_path):
        try:
            with open(crypto_jobs_path, 'r') as f:
                db_data = json.load(f).get('data', [])
                
                # Filter for companies that are enabled and category is 'crypto', 
                # but whose scraper type is NOT Greenhouse, Lever, Ashby, or Yzilabs.
                live_scrapers = ["GREENHOUSE", "LEVER", "ASHBYHQ", "YZILABS"]
                db_companies = {c.get('name').lower() for c in active_companies if c.get('scraper') not in live_scrapers}
                
                db_jobs_count = 0
                for dj in db_data:
                    comp = dj.get('company', '').lower()
                    if comp in db_companies:
                        jobs_pool.append(clean_job_data({
                            "company": dj.get('company'),
                            "title": dj.get('title'),
                            "location": dj.get('location', 'Remote'),
                            "link": dj.get('link'),
                            "notes": "Pre-scraped job from database."
                        }))
                        db_jobs_count += 1
                print(f"Merged {db_jobs_count} custom board jobs from database (Coinbase, Circle, Ripple, Gemini, etc.).")
        except Exception as e:
            print(f"Error loading pre-scraped jobs: {e}")

    # Save/update crypto_jobs.json with all active jobs fetched in this run
    try:
        crypto_jobs_path = os.path.join(base_dir, "crypto_jobs.json")
        existing_data = []
        if os.path.exists(crypto_jobs_path):
            with open(crypto_jobs_path, 'r', encoding='utf-8') as f:
                c = json.load(f)
                existing_data = c.get('data', []) if isinstance(c, dict) else c

        jobs_map = {}
        for item in existing_data:
            link = (item.get('link') or '').strip()
            comp = (item.get('company') or '').strip().lower()
            title = (item.get('title') or '').strip().lower()
            key = link if link else f"{comp}:{title}"
            if key:
                jobs_map[key] = item

        for item in jobs_pool:
            link = (item.get('link') or '').strip()
            comp = (item.get('company') or '').strip()
            title = (item.get('title') or '').strip()
            key = link if link else f"{comp.lower()}:{title.lower()}"
            if key and comp and title:
                jobs_map[key] = {
                    "company": comp,
                    "title": title,
                    "location": (item.get('location') or 'Remote').strip(),
                    "link": link
                }

        with open(crypto_jobs_path, 'w', encoding='utf-8') as f:
            json.dump({"data": list(jobs_map.values())}, f, indent=4)
        print(f"🎉 Updated crypto_jobs.json master database with {len(jobs_map)} total scraped jobs.")
    except Exception as e:
        print(f"Warning: Could not update crypto_jobs.json: {e}")

    return jobs_pool

def get_safe_sheet_name(cat_id, cat_name):
    # Excel sheet title limit is 31 characters
    clean_name = cat_name.replace(" / ", " & ").replace(" /", " & ").replace("/ ", " & ").replace("/", " & ")
    # Remove openpyxl forbidden characters: : ? * [ ] \ /
    for char in [":", "?", "*", "[", "]", "\\"]:
        clean_name = clean_name.replace(char, "")
    sheet_title = f"{cat_id:02d}. {clean_name}"
    if len(sheet_title) > 31:
        sheet_title = sheet_title[:28] + "..."
    return sheet_title

def export_master_excel(final_sheets):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    font_family = "Segoe UI"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10)
    data_font_bold = Font(name=font_family, size=10, bold=True)
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="E2E8F0")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    date_str = datetime.now().strftime("%d %b %Y")
    
    for cat_id, (cat_name, _) in categories.items():
        sheet_title = get_safe_sheet_name(cat_id, cat_name)
        ws = wb.create_sheet(title=sheet_title)
        
        # Ensure grid lines are visible
        ws.views.sheetView[0].showGridLines = True
        
        # Title Banner block
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"Web3 Jobs Research - {cat_name}"
        title_cell.font = Font(name=font_family, size=14, bold=True, color="1F4E78")
        title_cell.alignment = align_left
        ws.row_dimensions[1].height = 35
        
        # Spacer row
        ws.row_dimensions[2].height = 10
        
        # Header Row (Row 3)
        headers = ["#", "Company", "Role Title", "Location / Remote", "Direct Job Link", "Date Found", "Notes"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.value = h
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center if col_idx in [1, 6] else align_left
            cell.border = thin_border
        ws.row_dimensions[3].height = 25
        
        # Data Rows
        jobs = final_sheets[cat_id]
        for idx, job in enumerate(jobs, 1):
            row_idx = idx + 3
            comp = job.get('company', '').capitalize()
            title = job.get('title', '')
            loc = job.get('location', 'Remote')
            link = job.get('link', '')
            j_date = job.get('date_found', date_str)
            notes = job.get('notes', "Direct company job board link.")
            
            row_data = [idx, comp, title, loc, link, j_date, notes]
            row_fill = zebra_fill if idx % 2 == 0 else white_fill
            
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = val
                cell.font = data_font_bold if col_idx == 2 else data_font
                cell.fill = row_fill
                cell.border = thin_border
                
                # Alignments
                if col_idx == 1:
                    cell.alignment = align_center
                elif col_idx == 6:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                    
                # Hyperlink formatting
                if col_idx == 5 and val:
                    cell.value = val
                    cell.hyperlink = val
                    cell.font = Font(name=font_family, size=10, color="1F4E78", underline="single")
            
            ws.row_dimensions[row_idx].height = 20
            
        # Set manual widths
        column_widths = {
            "A": 6,   # #
            "B": 20,  # Company
            "C": 35,  # Role Title
            "D": 25,  # Location
            "E": 50,  # Direct Job Link (Raw URL)
            "F": 15,  # Date Found
            "G": 40   # Notes
        }
        for col_let, width in column_widths.items():
            ws.column_dimensions[col_let].width = width
            
    # Save Workbook
    excel_path = os.path.join(base_dir, "Web3_Jobs_Research_Master.xlsx")
    wb.save(excel_path)
    print(f"🎉 Master Excel Workbook successfully saved to: {excel_path}")

# Main runner pipeline
def run_pipeline():
    # Fetch all jobs
    raw_jobs = get_all_jobs()
    if not raw_jobs:
        print("No jobs fetched. Check internet connection and configuration.")
        return
        
    # Load posted jobs history
    posted_jobs_path = os.path.join(base_dir, "posted_jobs.json")
    posted_urls = set()
    posted_roles = set()  # Store normalized (company, title) tuples
    
    if os.path.exists(posted_jobs_path):
        try:
            with open(posted_jobs_path, 'r', encoding='utf-8') as f:
                history = json.load(f)
                posted_urls = set(history.get('urls', []))
                posted_roles = {tuple(x) for x in history.get('roles', [])}
        except Exception as e:
            print(f"Error loading posted jobs history: {e}")
        
    # Deduplicate by link AND company + title to avoid duplicate listings in same run
    deduped = []
    seen_links = set()
    seen_roles = set()
    
    for j in raw_jobs:
        link = j.get('link')
        company = j.get('company', '').strip().lower()
        title = j.get('title', '').strip().lower()
        
        # Normalize title (ignore case, brackets, remote indicators)
        norm_title = re.sub(r'\s*\([^)]*\)', '', title)
        norm_title = re.sub(r'\s*-\s*remote.*', '', norm_title)
        norm_title = norm_title.strip()
        
        role_key = (company, norm_title)
        
        if link and link not in seen_links and role_key not in seen_roles:
            # Check if this job has been posted in previous runs (history)
            is_relisted = (link in posted_urls or role_key in posted_roles)
            
            # Skip old/relisted jobs entirely to ensure only fresh/new jobs are published
            if not is_relisted:
                j['is_relisted'] = False
                deduped.append(j)
                seen_links.add(link)
                seen_roles.add(role_key)
            
    print(f"Deduplicated to {len(deduped)} unique positions.")
    
    # Categorize jobs with deduplication to the closest fitting category based on priority
    # Higher index in list = lower priority
    category_priority = [
        # Functional / Specific niches (highest priority)
        34, 10, 35, 2, 21, 20, 1, 3, 16, 5, 9, 19, 6, 22, 23, 24, 25, 33, 32, 18,
        # Specific locations
        4, 17, 8, 11, 26, 27, 28, 29, 30, 31, 7,
        # General categories (lowest priority)
        13, 14, 15
    ]
    
    categorized = {i: [] for i in categories}
    for j in deduped:
        title = j.get('title', '')
        location = j.get('location', '')
        company = j.get('company', '')
        if not title:
            continue
            
        matched_cats = []
        for cat_id, (cat_name, pattern) in categories.items():
            if cat_id == 12: # Signal Board is handled separately
                continue
            match = False
            # Check if this category is location-based
            if cat_id in [4, 7, 8, 11, 17, 26, 27, 28, 29, 30, 31]:
                if re.search(pattern, location.lower()) or re.search(pattern, title.lower()):
                    match = True
            elif cat_id == 10:
                if re.search(pattern, title.lower()) or re.search(pattern, company.lower()):
                    match = True
            else:
                if re.search(pattern, title.lower()):
                    match = True
            if match:
                # Exclude marketing or HR/talent roles from BD / Sales / Partnerships (category 1)
                if cat_id == 1:
                    if any(kw in title.lower() for kw in ["marketing", "hr ", " hr", "human resources", "recruiter", "talent", "people", "compensation", "business partner"]):
                        match = False
            if match:
                matched_cats.append(cat_id)
                
        if matched_cats:
            # Assign to the closest fitting category (highest priority)
            best_cat = min(matched_cats, key=lambda c: category_priority.index(c))
            categorized[best_cat].append(j)
            
            # Post best picks to Signal Board too (All Roles - Category 12)
            categorized[12].append(j)

    # Prioritize manual entries, then preferred companies, then the rest
    preferred = ['coinbase', 'circle', 'ripple', 'gemini', 'paxos', 'worldcoin', 'sygnum', 'consensys', 'monad', 'aster', 'nym', 'vibe', 'manifolds.ai', 'arbitrum', 'arbitrumfoundation', 'gsr', 'b2c2', 'hudsonrivertrading', 'portofinotechnologies', 'drw', 'keyrock', 'wintermute', 'jumpcrypto']
    def job_priority(j):
        if j.get('is_manual'):
            return 0
        if j.get('company', '').lower() in preferred:
            return 1
        return 2

    # Sort each categorized list by priority
    for cat_id in categorized:
        categorized[cat_id].sort(key=job_priority)

    # Validate links and select top 5 unique jobs
    print("Selecting 5 unique-company roles per category sheet...")
    
    # Identify links to verify (top 15 jobs in sorted order for each category to cover fallbacks)
    links_to_check = set()
    for cat_id in categories:
        for job in categorized[cat_id][:15]:
            link = job.get('link')
            if link:
                links_to_check.add(link)
    links_to_check = list(links_to_check)

    # Parallel URL check
    print(f"Verifying {len(links_to_check)} candidate links...")
    verified = {}
    with ThreadPoolExecutor(max_workers=20) as executor:
        future_to_url = {executor.submit(verify_url, url): url for url in links_to_check}
        for future in as_completed(future_to_url):
            url = future_to_url[future]
            try:
                verified[url] = future.result()
            except Exception:
                verified[url] = False

    # Apply verifications and enforce unique companies
    date_str = datetime.now().strftime("%d %b %Y")
    final_sheets = {}
    
    for cat_id, (cat_name, _) in categories.items():
        cat_jobs = categorized[cat_id]
        
        final_list = []
        seen_companies = set()
        
        # Enforce uniqueness + verified links
        for job in cat_jobs:
            company = job.get('company', '').lower()
            link = job.get('link')
            
            if company in seen_companies:
                continue
                
            is_working = verified.get(link, False)
            if is_working:
                final_list.append(job)
                seen_companies.add(company)
                if len(final_list) == 5:
                    break
                    
        # Relax uniqueness if under 5
        if len(final_list) < 5:
            for job in cat_jobs:
                if job not in final_list:
                    link = job.get('link')
                    is_working = verified.get(link, False)
                    if is_working:
                        final_list.append(job)
                        if len(final_list) == 5:
                            break
                            
        # Final fallback whatever we have (only if they are verified as working)
        if len(final_list) < 5:
            for job in cat_jobs:
                if job not in final_list:
                    link = job.get('link')
                    is_working = verified.get(link, False)
                    if is_working:
                        final_list.append(job)
                        if len(final_list) == 5:
                            break
                        
        final_sheets[cat_id] = final_list[:5]

    # Export CSVs
    os.makedirs(output_dir, exist_ok=True)
    for cat_id, (cat_name, _) in categories.items():
        safe_name = cat_name.replace("/", "_").replace("&", "_").replace(" ", "_").replace("__", "_").strip("_")
        filepath = os.path.join(output_dir, f"{cat_id:02d}_{safe_name}.csv")
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write("#,Company,Role Title,Location / Remote,Direct Job Link,Date Found,Notes\n")
            for idx, job in enumerate(final_sheets[cat_id]):
                comp = job.get('company', '').capitalize()
                title = job.get('title', '').replace(',', ' ').replace('\n', ' ')
                loc = job.get('location', 'Remote').replace(',', ' ').replace('\n', ' ')
                link = job.get('link', '')
                j_date = job.get('date_found', date_str)
                notes = job.get('notes', "Direct company job board link.")
                f.write(f"{idx+1},{comp},{title},{loc},{link},{j_date},{notes}\n")

    # Export Master Excel Workbook
    export_master_excel(final_sheets)

    # Update posted jobs history with newly selected jobs
    for cat_id, jobs in final_sheets.items():
        for job in jobs:
            link = job.get('link')
            company = job.get('company', '').strip().lower()
            title = job.get('title', '').strip().lower()
            
            norm_title = re.sub(r'\s*\([^)]*\)', '', title)
            norm_title = re.sub(r'\s*-\s*remote.*', '', norm_title)
            norm_title = norm_title.strip()
            
            if link:
                posted_urls.add(link)
            if company and norm_title:
                posted_roles.add((company, norm_title))
                
    try:
        with open(posted_jobs_path, 'w', encoding='utf-8') as f:
            json.dump({
                "urls": list(posted_urls),
                "roles": [list(r) for r in posted_roles]
            }, f, indent=4)
        print("🎉 Updated posted jobs history database.")
    except Exception as e:
        print(f"Error saving posted jobs history: {e}")

    # Generate HTML Dashboard
    js_data = []
    for cat_id, (cat_name, _) in categories.items():
        cat_rows = []
        for idx, job in enumerate(final_sheets[cat_id]):
            cat_rows.append({
                "num": idx + 1,
                "company": job.get('company', '').capitalize(),
                "title": job.get('title', ''),
                "location": job.get('location', 'Remote').replace('\n', ', '),
                "link": job.get('link', ''),
                "date": job.get('date_found', date_str),
                "notes": job.get('notes', "Direct company job board link."),
                "is_relisted": job.get('is_relisted', False)
            })
        js_data.append({
            "num": cat_id,
            "category": cat_name,
            "channel": telegram_channels[cat_id],
            "jobs": cat_rows
        })

    write_dashboard_html(js_data)
    write_dashboard_ts(js_data)
    
    # Automatically sync scraped jobs to the Web3 Job Seeker portal dataset
    sync_script_path = os.path.join(base_dir, "scripts", "sync_scraped_jobs.py")
    if os.path.exists(sync_script_path):
        try:
            import subprocess
            print("🚀 Syncing scraped jobs with Web3 Job Seeker portal database...")
            subprocess.run(["python3", sync_script_path], check=False)
        except Exception as e:
            print(f"⚠️ Warning: Auto-sync script failed: {e}")

    print(f"🎉 CSV Spreadsheets successfully exported to {output_dir}/")
    print(f"🎉 Interactive Dashboard updated at {dashboard_path}")

def write_dashboard_html(js_data):
    # Same interactive dashboard template from earlier
    date_str = datetime.now().strftime("%d %b %Y")
    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Interactive Web3 Jobs Dashboard</title>
<link href="https://fonts.googleapis.com/css2?family=Outfit:wght@400;500;600;700&family=JetBrains+Mono&display=swap" rel="stylesheet">
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    font-family: 'Outfit', Arial, sans-serif;
    background: #0b0d13;
    color: #e2e8f0;
    min-height: 100vh;
    padding-bottom: 60px;
  }}

  .header {{
    background: linear-gradient(135deg, #161a29 0%, #0b0d13 100%);
    border-bottom: 1px solid #1f293d;
    padding: 40px 60px;
    text-align: center;
    position: relative;
    overflow: hidden;
  }}
  
  .header::before {{
    content: '';
    position: absolute;
    top: -50%;
    left: -50%;
    width: 200%;
    height: 200%;
    background: radial-gradient(circle, rgba(99, 102, 241, 0.1) 0%, transparent 70%);
    pointer-events: none;
  }}

  .header h1 {{
    font-size: 36px;
    font-weight: 700;
    color: #fff;
    margin-bottom: 12px;
    letter-spacing: -0.5px;
    background: linear-gradient(90deg, #6366f1, #a855f7);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
  }}
  
  .header p {{
    color: #94a3b8;
    font-size: 16px;
    max-width: 600px;
    margin: 0 auto;
    line-height: 1.5;
  }}

  .container {{
    max-width: 1500px;
    margin: 0 auto;
    padding: 40px;
  }}

  .category-select-wrapper {{
    background: #141824;
    border: 1px solid #1f293d;
    border-radius: 16px;
    padding: 24px;
    margin-bottom: 32px;
    box-shadow: 0 4px 20px rgba(0, 0, 0, 0.2);
  }}

  .category-label {{
    font-size: 14px;
    font-weight: 600;
    color: #818cf8;
    margin-bottom: 12px;
    text-transform: uppercase;
    letter-spacing: 1px;
  }}

  .category-grid {{
    display: grid;
    grid-template-columns: repeat(auto-fill, minmax(280px, 1fr));
    gap: 10px;
  }}

  .category-btn {{
    background: #1e2538;
    border: 1px solid #2e3b56;
    color: #94a3b8;
    padding: 12px 16px;
    border-radius: 10px;
    font-size: 13.5px;
    font-weight: 500;
    text-align: left;
    cursor: pointer;
    transition: all 0.2s ease;
    white-space: nowrap;
    overflow: hidden;
    text-overflow: ellipsis;
  }}

  .category-btn:hover {{
    background: #252e46;
    border-color: #435375;
    color: #fff;
    transform: translateY(-1px);
  }}

  .category-btn.active {{
    background: linear-gradient(135deg, #4f46e5 0%, #3730a3 100%);
    border-color: #6366f1;
    color: #fff;
    box-shadow: 0 4px 12px rgba(79, 70, 229, 0.3);
  }}

  /* Active Category Card */
  .active-card {{
    background: #141824;
    border: 1px solid #1f293d;
    border-radius: 16px;
    overflow: hidden;
    box-shadow: 0 4px 30px rgba(0, 0, 0, 0.3);
    margin-bottom: 30px;
  }}

  .card-header {{
    background: #1a2032;
    padding: 24px 32px;
    border-bottom: 1px solid #2e3b56;
    display: flex;
    justify-content: space-between;
    align-items: center;
    flex-wrap: wrap;
    gap: 16px;
  }}

  .card-title-group {{
    display: flex;
    flex-direction: column;
    gap: 4px;
  }}

  .card-title {{
    font-size: 22px;
    font-weight: 700;
    color: #fff;
  }}

  .card-subtitle {{
    font-size: 13px;
    color: #6366f1;
    font-family: 'JetBrains Mono', monospace;
  }}

  .channel-badge {{
    background: rgba(99, 102, 241, 0.15);
    border: 1px solid rgba(99, 102, 241, 0.3);
    color: #818cf8;
    padding: 6px 14px;
    border-radius: 30px;
    font-size: 13px;
    font-weight: 600;
    text-decoration: none;
    transition: all 0.2s;
    display: inline-flex;
    align-items: center;
    gap: 6px;
  }}

  .channel-badge:hover {{
    background: rgba(99, 102, 241, 0.25);
    border-color: #6366f1;
    color: #fff;
  }}

  /* Table styles */
  .table-wrapper {{
    padding: 32px;
    overflow-x: auto;
  }}

  .job-table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 14px;
    text-align: left;
  }}

  .job-table th {{
    background: #1c2235;
    color: #94a3b8;
    font-weight: 600;
    padding: 16px 20px;
    font-size: 12px;
    letter-spacing: 0.8px;
    text-transform: uppercase;
    border-bottom: 2px solid #2e3b56;
  }}

  .job-table td {{
    padding: 16px 20px;
    border-bottom: 1px solid #1f293d;
    color: #cbd5e1;
    vertical-align: middle;
  }}

  .job-table tr:hover td {{
    background: #181f30;
    color: #fff;
  }}

  .num-col {{
    width: 60px;
    font-family: 'JetBrains Mono', monospace;
    color: #64748b;
    font-weight: 600;
  }}

  .company-col {{
    font-weight: 600;
    color: #fff;
  }}

  .link-col a {{
    color: #60a5fa;
    text-decoration: none;
    font-size: 13px;
    transition: color 0.15s;
    word-break: break-all;
  }}

  .link-col a:hover {{
    color: #93c5fd;
    text-decoration: underline;
  }}

  .actions {{
    display: flex;
    justify-content: flex-end;
    padding: 0 32px 32px;
    gap: 16px;
  }}

  .btn {{
    background: #22c55e;
    color: #fff;
    border: none;
    border-radius: 8px;
    padding: 12px 24px;
    font-size: 14px;
    font-weight: 600;
    cursor: pointer;
    display: inline-flex;
    align-items: center;
    gap: 8px;
    transition: all 0.15s;
  }}

  .btn:hover {{
    background: #16a34a;
    transform: translateY(-1px);
  }}

  .btn-copy-tsv {{
    background: #3b82f6;
  }}

  .btn-copy-tsv:hover {{
    background: #2563eb;
  }}

  .btn.copied {{
    background: #059669;
  }}

  .notes-badge {{
    background: #1e293b;
    border: 1px solid #334155;
    padding: 2px 8px;
    border-radius: 4px;
    font-size: 11px;
    color: #94a3b8;
  }}

  @media (max-width: 1024px) {{
    .container {{ padding: 20px; }}
    .header {{ padding: 30px 20px; }}
  }}
</style>
</head>
<body>

<div class="header">
  <h1>🌐 Web3 Job Research Dashboard</h1>
  <p>Systematically compiled spreadsheet datasets containing 5 verified direct-link roles per job category</p>
  <div style="margin-top: 20px;">
    <a href="Web3_Jobs_Research_Master.xlsx" class="btn" style="background: linear-gradient(135deg, #107c41 0%, #0b5930 100%); border: 1px solid #107c41; box-shadow: 0 4px 15px rgba(16, 124, 65, 0.4); text-decoration: none;" download>
      📥 Download Master Excel Workbook (.xlsx)
    </a>
  </div>
</div>

<div class="container">
  
  <!-- Category Selector -->
  <div class="category-select-wrapper">
    <div class="category-label">Select Job Category Sheets</div>
    <div class="category-grid" id="categoryGrid"></div>
  </div>

  <!-- Table Card -->
  <div class="active-card">
    <div class="card-header">
      <div class="card-title-group">
        <div class="card-title" id="activeCategoryName">BD / Sales / Partnerships</div>
        <div class="card-subtitle" id="activeCategoryMetadata">Category #1 • CSV Format Ready</div>
      </div>
      <a href="#" class="channel-badge" id="activeChannelBadge" target="_blank">
        ✈️ Join Channel
      </a>
    </div>

    <div class="table-wrapper">
      <table class="job-table">
        <thead>
          <tr>
            <th class="num-col">#</th>
            <th>Company</th>
            <th>Role Title</th>
            <th>Location / Remote</th>
            <th>Direct Job Link</th>
            <th>Date Found</th>
            <th>Notes</th>
          </tr>
        </thead>
        <tbody id="jobTableBody"></tbody>
      </table>
    </div>

    <div class="actions">
      <button class="btn btn-copy-tsv" onclick="copyTSV()" id="copyTsvBtn">
        📋 Copy Sheet to Clipboard (TSV)
      </button>
      <button class="btn" onclick="copyTelegramText()" id="copyTgBtn">
        ✈️ Copy Telegram Post Text
      </button>
    </div>
  </div>

</div>

<script>
// Seed Data
const dashboardData = {json.dumps(js_data, indent=2)};

let currentCategoryId = 1;

function init() {{
  const grid = document.getElementById('categoryGrid');
  grid.innerHTML = '';
  
  dashboardData.forEach(cat => {{
    const btn = document.createElement('button');
    btn.className = `category-btn ${{cat.num === currentCategoryId ? 'active' : ''}}`;
    btn.id = `cat-btn-${{cat.num}}`;
    btn.textContent = `${{cat.num}}. ${{cat.category}}`;
    btn.onclick = () => selectCategory(cat.num);
    grid.appendChild(btn);
  }});
  
  renderTable();
}}

function selectCategory(id) {{
  document.getElementById(`cat-btn-${{currentCategoryId}}`).classList.remove('active');
  currentCategoryId = id;
  document.getElementById(`cat-btn-${{currentCategoryId}}`).classList.add('active');
  renderTable();
}}

function renderTable() {{
  const cat = dashboardData.find(c => c.num === currentCategoryId);
  if (!cat) return;
  
  document.getElementById('activeCategoryName').textContent = cat.category;
  document.getElementById('activeCategoryMetadata').textContent = `Category #${{cat.num}} • CSV Format Ready`;
  
  const badge = document.getElementById('activeChannelBadge');
  badge.href = `https://${{cat.channel}}`;
  badge.textContent = `✈️ ${{cat.channel}}`;
  
  const tbody = document.getElementById('jobTableBody');
  tbody.innerHTML = '';
  
  cat.jobs.forEach(job => {{
    const tr = document.createElement('tr');
    tr.innerHTML = `
      <td class="num-col">${{job.num}}</td>
      <td class="company-col">${{job.company}}</td>
      <td style="font-weight: 500;">${{job.title}}</td>
      <td>${{job.location}}</td>
      <td class="link-col"><a href="${{job.link}}" target="_blank">${{job.link}}</a></td>
      <td style="font-family: 'JetBrains Mono', monospace; font-size:12px; color:#94a3b8;">${{job.date}}</td>
      <td><span class="notes-badge">${{job.notes}}</span></td>
    `;
    tbody.appendChild(tr);
  }});
}}

function copyTSV() {{
  const cat = dashboardData.find(c => c.num === currentCategoryId);
  if (!cat) return;
  
  const headers = ["#", "Company", "Role Title", "Location / Remote", "Direct Job Link", "Date Found", "Notes"];
  const rows = cat.jobs.map(j => [
    j.num,
    j.company,
    j.title.replace(/\\t/g, ' '),
    j.location.replace(/\\t/g, ' '),
    j.link,
    j.date,
    j.notes
  ]);
  
  const tsvContent = [headers, ...rows].map(r => r.join('\\t')).join('\\n');
  
  navigator.clipboard.writeText(tsvContent).then(() => {{
    const btn = document.getElementById('copyTsvBtn');
    btn.textContent = '✅ Copied TSV to Clipboard!';
    btn.classList.add('copied');
    setTimeout(() => {{
      btn.textContent = '📋 Copy Sheet to Clipboard (TSV)';
      btn.classList.remove('copied');
    }}, 2000);
  }});
}}

function copyTelegramText() {{
  const cat = dashboardData.find(c => c.num === currentCategoryId);
  if (!cat) return;
  
  let text = `📢 **New Web3 Job Listings: ${{cat.category}}**\\n\\n`;
  cat.jobs.forEach(j => {{
    text += `💼 **${{j.title}}** at **${{j.company}}**\\n📍 ${{j.location}}\\n🔗 Direct Apply: ${{j.link}}\\n\\n`;
  }});
  text += `Join for more: https://${{cat.channel}}`;
  
  navigator.clipboard.writeText(text).then(() => {{
    const btn = document.getElementById('copyTgBtn');
    btn.textContent = '✅ Copied Post to Clipboard!';
    btn.classList.add('copied');
    setTimeout(() => {{
      btn.textContent = '✈️ Copy Telegram Post Text';
      btn.classList.remove('copied');
    }}, 2000);
  }});
}}

init();
</script>
</body>
</html>
"""
    with open(dashboard_path, 'w', encoding='utf-8') as f:
        f.write(html_content)

def write_dashboard_ts(js_data):
    react_dashboard_data_path = os.path.join(base_dir, "frontend", "src", "data", "dashboard.ts")
    date_str = datetime.now().strftime("%d %b %Y")
    
    ts_content = f"""export type Job = {{
  num: number;
  company: string;
  title: string;
  location: string;
  link: string;
  date: string;
  notes: string;
  is_relisted?: boolean;
}};

export type Category = {{
  num: number;
  category: string;
  channel: string;
  jobs: Job[];
}};

export const LAST_UPDATED = "{date_str}";

export const dashboardData: Category[] = {json.dumps(js_data, indent=2)};

export function telegramUrl(channel: string): string {{
  const handle = channel.replace(/^https?:\\/\\//, "").replace(/^t\\.me\\//, "");
  return `https://t.me/${{handle}}`;
}}

export function totalJobs(categories: Category[]): number {{
  return categories.reduce((sum, item) => sum + item.jobs.length, 0);
}}
"""
    os.makedirs(os.path.dirname(react_dashboard_data_path), exist_ok=True)
    with open(react_dashboard_data_path, 'w', encoding='utf-8') as f:
        f.write(ts_content)

# CLI parser
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Web3 Job Research Automation Tool")
    parser.add_argument("--run", action="store_true", help="Run the automated search & compilation pipeline")
    parser.add_argument("--add", action="store_true", help="Interactively add a manual job entry for custom boards")
    parser.add_argument("--stats", action="store_true", help="Show current manual and database counts")
    
    args = parser.parse_args()
    
    if args.add:
        add_manual_job()
    elif args.stats:
        manual_jobs = []
        if os.path.exists(manual_jobs_path):
            with open(manual_jobs_path, 'r') as f:
                manual_jobs = json.load(f)
        print(f"Total manual custom board entries: {len(manual_jobs)}")
    else:
        # Default is running the pipeline
        run_pipeline()
